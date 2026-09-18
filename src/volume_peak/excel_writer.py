"""신고 거래량 종목 목록을 엑셀(.xlsx)로 저장한다."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .screener import Breakout

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
BODY_FONT = Font(size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

INT_FORMAT = "#,##0"
RATE_FORMAT = "+0.00;-0.00;0.00"
RATIO_FORMAT = '#,##0.00"배"'

# (헤더, 필드, 너비, 정렬, 표시형식)
COLUMNS: list[tuple[str, str, int, str, str]] = [
    ("번호", "no", 6, "center", ""),
    ("시장", "market", 10, "center", ""),
    ("종목코드", "code", 10, "center", "@"),
    ("종목명", "name", 28, "left", ""),
    ("기준일 거래량", "volume", 16, "right", INT_FORMAT),
    ("직전 최고거래량", "prev_peak", 16, "right", INT_FORMAT),
    ("직전 기록일", "prev_peak_date", 13, "center", ""),
    ("배수", "ratio", 9, "right", RATIO_FORMAT),
    ("종가", "close", 11, "right", INT_FORMAT),
    ("등락률(%)", "change_rate", 11, "right", RATE_FORMAT),
    ("거래대금(원)", "value", 18, "right", INT_FORMAT),
    ("이력 시작일", "history_from", 13, "center", ""),
    ("비고", "note", 24, "left", ""),
]


@dataclass
class ExcelResult:
    """엑셀 생성 결과."""

    path: Path
    row_count: int


def _sheet_title(as_of: str) -> str:
    return f"신고거래량_{as_of.replace('-', '')}"[:31]


def _cell_value(item: Breakout, field: str, index: int):
    if field == "no":
        return index
    if field == "ratio":
        return item.ratio
    return getattr(item, field, "")


def write_excel(
    breakouts: list[Breakout],
    out_path: str | Path,
    as_of: str,
    period: str = "",
    source_note: str = "출처: 한국거래소 정보데이터시스템(data.krx.co.kr) 전종목 시세 · 개별종목 시세 추이",
) -> ExcelResult:
    """신고 거래량 종목을 서식이 적용된 엑셀로 저장한다."""
    path = Path(out_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _sheet_title(as_of)

    scope = f", {period}" if period else ""
    title = f"신고 거래량 종목 ({as_of} 기준{scope}, 총 {len(breakouts)}종목)"
    sheet.cell(row=1, column=1, value=title).font = Font(size=13, bold=True)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    sheet.cell(row=2, column=1, value=source_note).font = Font(size=9, color="808080")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(COLUMNS))

    header_row = 4
    for index, (header, _field, width, _align, _fmt) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=header_row, column=index, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[header_row].height = 24

    for offset, item in enumerate(breakouts):
        excel_row = header_row + 1 + offset
        for index, (_header, field, _width, align, number_format) in enumerate(COLUMNS, start=1):
            value = _cell_value(item, field, offset + 1)
            cell = sheet.cell(row=excel_row, column=index, value=value)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(horizontal=align, vertical="center")
            if number_format:
                # 종목코드는 앞자리 0 이 사라지지 않도록 텍스트 서식으로 둔다.
                cell.number_format = number_format

    last_row = header_row + len(breakouts)
    if breakouts:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(COLUMNS))}{last_row}"
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    workbook.save(path)
    return ExcelResult(path=path, row_count=len(breakouts))
